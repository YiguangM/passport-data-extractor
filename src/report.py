"""
report.py
---------
Writes a list of PassportRecord objects out to a styled Excel workbook:
  - "Passports" sheet: one row per file, with a status color.
  - "Summary" sheet: run stats (total processed, OK/partial/failed counts,
    how many needed OCR, how many passed MRZ checksum validation).
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from extractor import PassportRecord

HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)

STATUS_FILLS = {
    "ok": PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid"),
    "partial": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
    "failed": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
}

COLUMNS = [
    ("file_name", "File Name", 28),
    ("given_names", "Given Names", 20),
    ("surname", "Surname", 18),
    ("sex", "Sex", 10),
    ("date_of_birth", "Date of Birth", 14),
    ("nationality", "Nationality", 14),
    ("passport_number", "Passport #", 16),
    ("issuing_country", "Issuing Country", 14),
    ("date_of_expiry", "Date of Expiry", 14),
    ("personal_number", "Personal #", 16),
    ("mrz_found", "MRZ Found", 10),
    ("mrz_valid", "MRZ Valid", 10),
    ("extraction_method", "Extraction Method", 16),
    ("status", "Status", 10),
    ("notes", "Notes", 40),
]


def write_report(records: list[PassportRecord], output_path: Path) -> Path:
    wb = Workbook()

    _write_passports_sheet(wb.active, records)
    _write_summary_sheet(wb.create_sheet("Summary"), records)

    wb.save(output_path)
    return output_path


def _write_passports_sheet(ws, records: list[PassportRecord]) -> None:
    ws.title = "Passports"

    for col_idx, (_, header, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"

    for row_idx, record in enumerate(records, start=2):
        row_fill = STATUS_FILLS.get(record.status)
        for col_idx, (field_name, _, _) in enumerate(COLUMNS, start=1):
            value = getattr(record, field_name)
            if field_name in ("used_ocr", "mrz_found"):
                value = "Yes" if value else "No"
            elif field_name == "mrz_valid":
                value = {True: "Yes", False: "No", None: "N/A"}[value]
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if row_fill:
                cell.fill = row_fill


def _write_summary_sheet(ws, records: list[PassportRecord]) -> None:
    total = len(records)
    ok = sum(1 for r in records if r.status == "ok")
    partial = sum(1 for r in records if r.status == "partial")
    failed = sum(1 for r in records if r.status == "failed")
    ocr_used = sum(1 for r in records if r.used_ocr)
    mrz_valid = sum(1 for r in records if r.mrz_valid)

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 20

    rows = [
        ("Run generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Total files processed", total),
        ("Fully extracted (OK)", ok),
        ("Partial extraction", partial),
        ("Failed", failed),
        ("Files requiring OCR", ocr_used),
        ("MRZ checksum valid", mrz_valid),
    ]
    for r_idx, (label, value) in enumerate(rows, start=1):
        label_cell = ws.cell(row=r_idx, column=1, value=label)
        label_cell.font = Font(bold=True)
        ws.cell(row=r_idx, column=2, value=value)
